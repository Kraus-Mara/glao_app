# Copyright (c) 2026, kr and contributors
# For license information, please see license.txt

import frappe
from frappe.model.document import Document
from frappe.utils import now
from frappe.model.naming import make_autoname
from frappe.utils.pdf import get_pdf


class Expedition(Document):
	# begin: auto-generated types
	# This code is auto-generated. Do not modify anything in this block.

	from typing import TYPE_CHECKING

	if TYPE_CHECKING:
		from frappe.types import DF

		carrier: DF.Data | None
		client: DF.Data | None
		delivery_date: DF.Date | None
		dmc: DF.Link | None
		expedition_date: DF.Date | None
		job_no: DF.Data | None
		project: DF.Link | None
		status: DF.Literal["Waiting", "Shipped"]
	# end: auto-generated types

	def autoname(self):
		self.name = make_autoname(str(self.project) + "-" + str(self.client) + " expedition " + ".#")

	def validate(self):
		previous = self.get_doc_before_save()
		previous_status = previous.status if previous else None

		if self.status == "Shipped" and previous_status != "Shipped":
			if not self.expedition_date:
				self.expedition_date = now()
			self._send_dmc()
		else:
			return 1

	def _send_dmc(self):
		"""Transfère les articles vers CLIENTS/SITE et décrémente le stock réservé."""
		dmc_items = frappe.get_all("Gestion DMC Items", filters=[["parent", "=", self.dmc]])
		dmc = frappe.get_doc("Gestion DMC", str(self.dmc))
		project = frappe.get_doc("Projects", self.project)
		for doc in dmc_items:
			r = frappe.get_doc("Gestion DMC Items", doc.name)
			if r.no_serving:
				continue
			if r.reserved:
				frappe.new_doc(
					doctype="Movement",
					type="Transfert",
					article_from_stock=r.item_from_stock,
					quantity_to_manipulate=r.true_quantity,
					source_place=r.source_place,
					target_place="CLIENTS/" + str(dmc.client) + "/" + str(dmc.project),
				).save()

				sd = frappe.get_doc("Stock", str(r.item_from_stock), for_update=True)
				sd.reserved_quantity -= r.true_quantity
				sd.save(ignore_permissions=True)
				if project:
					project.append(
						"project_items_sent",
						{
							"article": r.item_from_stock,
							"designation": r.get("designation"),
							"quantity": r.true_quantity,
							"source_place": r.source_place,
							"date": self.expedition_date,
							"expedition": self.name,
						},
					)
		dmc_compos = frappe.get_all("Gestion DMC Compositions", filters=[["parent", "=", self.dmc]])
		for d in dmc_compos:
			c = frappe.get_doc("Gestion DMC Compositions", d.name)
			if c.comp_saved:
				if not frappe.db.get_value("Composition", c.composition, "by_dmc", c.name):
					continue
				frappe.db.set_value("Composition", c.composition, "not_available", 1)
				frappe.db.set_value(
					"Composition",
					c.composition,
					"place",
					"CLIENTS/" + str(dmc.client) + "/" + str(dmc.project),
				)
				frappe.db.set_value("Composition", c.composition, "reserved", 0)
				frappe.db.set_value("Composition", c.composition, "by_dmc", None)
				if project:
					doc_comp = frappe.get_doc("Composition", c.composition)
					for r in doc_comp.items:
						project.append(
							"project_compositions_sent",
							{
								"composition": c.composition,
								"article": r.item,
								"designation": r.designation,
								"quantity": r.quantity,
								"source_place": r.saved_place,
								"date": self.expedition_date,
								"expedition": self.name,
							},
						)
		if project:
			project.save(ignore_permissions=True)

		frappe.db.set_value("Gestion DMC", dmc.name, "status", "Shipped")
		frappe.msgprint("Les articles et compositions ont été expédiés avec succès")


# ============================================================
# HELPERS
# ============================================================
@frappe.whitelist()
def _get_item_weight(item_code):
	"""Retourne le poids (float, kg) d'un article depuis ses caractéristiques."""
	if not item_code:
		return 0.0
	try:
		article = frappe.get_doc("Article", item_code)
	except frappe.DoesNotExistError:
		return 0.0

	total_kg = 0.0
	for row in article.chars or []:
		if (row.characteristics_type or "").upper() == "POIDS":
			try:
				value = float(row.value or 0)
			except TypeError, ValueError:
				value = 0.0
			unit = (row.unit or "").upper()
			if unit == "KG":
				total_kg += value
			elif unit == "G":
				total_kg += value / 1000.0
			elif unit == "T":
				total_kg += value * 1000.0
	return total_kg


@frappe.whitelist()
def _get_composition_weight(composition_name):
	"""Retourne le poids total (float, kg) d'une composition = somme des poids de ses articles."""
	if not composition_name:
		return 0.0
	try:
		composition = frappe.get_doc("Composition", composition_name)
	except frappe.DoesNotExistError:
		return 0.0

	total_kg = 0.0
	for row in composition.items or []:
		unit_weight = _get_item_weight(row.item)
		qty = float(row.quantity or 0)
		total_kg += unit_weight * qty
	return total_kg


# ============================================================
# COLLECTE DES DONNÉES
# ============================================================
def _collect_expedition_data(expedition):
	"""Collecte les articles, compositions et poids calculés pour une expédition donnée."""
	project = frappe.get_doc("Projects", expedition.project)
	dmc = frappe.get_doc("Gestion DMC", expedition.dmc)

	# --- Contact ---
	cdmc = frappe.get_all(
		"Creation DMC",
		filters=[["project", "=", str(expedition.project)], ["contact", "Is", "Set"]],
	)
	contact = "N/A"
	if cdmc:
		crdmc = frappe.get_doc("Creation DMC", cdmc[0])
		contact = crdmc.contact or "N/A"

	# --- Articles (uniquement quantité > 0) ---
	items_rows = []
	raw_items = frappe.get_all(
		"Gestion DMC Items",
		filters=[
			["parent", "=", expedition.dmc],
			["parenttype", "=", "Gestion DMC"],
			["no_serving", "=", 0],
			["true_quantity", ">", 0],
		],
		fields=["item_from_stock", "designation", "true_quantity", "closest_event"],
	)
	raw_items = [r for r in raw_items if float(r.get("true_quantity") or 0) > 0]

	total_weight_items = 0.0
	for row in raw_items:
		unit_weight = _get_item_weight(row.get("item_from_stock"))
		qty = float(row.get("true_quantity") or 0)
		line_weight = unit_weight * qty
		row["unit_weight"] = round(unit_weight, 3)
		row["line_weight"] = round(line_weight, 3)
		total_weight_items += line_weight
		items_rows.append(row)

	# --- Compositions (uniquement quantité > 0) ---
	bom_rows = []
	raw_boms = frappe.get_all(
		"Gestion DMC Compositions",
		filters=[
			["parent", "=", expedition.dmc],
			["parenttype", "=", "Gestion DMC"],
			["no_serving", "=", 0],
			["composition", "is", "Set"],
			["quantity", ">", 0],
		],
		fields=["composition", "quantity"],
	)
	raw_boms = [r for r in raw_boms if float(r.get("quantity") or 0) > 0]

	total_weight_bom = 0.0
	for row in raw_boms:
		qty = float(row.get("quantity") or 0)
		unit_weight = _get_composition_weight(row.get("composition"))
		line_weight = unit_weight * qty
		row["unit_weight"] = round(unit_weight, 3)
		row["line_weight"] = round(line_weight, 3)
		total_weight_bom += line_weight
		bom_rows.append(row)

	total_weight = round(total_weight_items + total_weight_bom, 3)

	return {
		"project": project,
		"dmc": dmc,
		"contact": contact,
		"items_rows": items_rows,
		"bom_rows": bom_rows,
		"total_weight_items": round(total_weight_items, 3),
		"total_weight_bom": round(total_weight_bom, 3),
		"total_weight": total_weight,
	}


# ============================================================
# EXPORT PDF
# ============================================================
@frappe.whitelist()
def export_expedition_pdf(name):
	expedition = frappe.get_doc("Expedition", name)
	carrier = getattr(expedition, "carrier", None)
	data = _collect_expedition_data(expedition)

	html_content = frappe.render_template(
		"""
		<style>
			@page { size: A4; margin: 8mm; }
			.bl-container { font-family: sans-serif; font-size: 8px; color: #222; }
			.bl-container * { font-size: 8px; box-sizing: border-box; }

			/* ------------------------------------------------------------
			   TABLE PRINCIPALE : le <thead> se répète sur chaque page
			   ------------------------------------------------------------ */
			.bl-main-table {
				width: 100%;
				border-collapse: collapse;
			}
			.bl-main-table > thead {
				display: table-header-group;
			}
			.bl-main-table > tbody > tr > td {
				padding: 0;
			}

			/* ---------- EN-TÊTE (compact) ---------- */
			.bl-header {
				border-bottom: 1.5px solid #003a70;
				padding-bottom: 3px;
				margin-bottom: 4px;
			}
			.bl-header table { width: 100%; border-collapse: collapse; }
			.bl-header h1 { color: #003a70; font-size: 10px; margin: 0 0 1px 0; font-weight: bold; }
			.bl-header .bl-num { font-size: 8px; margin: 0; }
			.bl-header .bl-right-block { text-align: right; }
			.bl-header .bl-right-block p { margin: 0; font-size: 8px; }

			/* ---------- EXPÉDITEUR / DESTINATAIRE (compact) ---------- */
			.bl-parties { width: 100%; border-collapse: collapse; margin-bottom: 4px; }
			.bl-parties td { vertical-align: top; padding: 3px 4px; }
			.bl-party { width: 50%; border: 1px solid #ccc; }
			.bl-party h3 { margin: 0 0 2px 0; font-size: 8px; color: #003a70; text-transform: uppercase; font-weight: bold; }
			.bl-party p { margin: 0; font-size: 8px; line-height: 1.15; }

			/* ---------- TRANSPORTEUR (compact) ---------- */
			.bl-transport { width: 100%; border-collapse: collapse; margin-bottom: 4px; }
			.bl-transport td { padding: 2px 4px; border: 1px solid #ccc; background: #f7f9fc; font-size: 8px; }

			/* ---------- SECTIONS ---------- */
			.bl-section {
				color: #003a70;
				font-size: 9px;
				font-weight: bold;
				margin: 6px 0 2px 0;
				border-bottom: 1px solid #003a70;
				padding-bottom: 1px;
			}

			/* ---------- TABLEAUX DE DONNÉES (police 8px) ---------- */
			.bl-table { width: 100%; border-collapse: collapse; margin-bottom: 4px; }
			.bl-table th, .bl-table td {
				border: 0.5px solid #bbb;
				padding: 1px 3px;
				font-size: 8px;
				line-height: 1.2;
			}
			.bl-table th {
				background: #003a70;
				color: #fff;
				font-weight: bold;
				text-align: left;
				font-size: 8px;
			}
			.bl-table tbody tr:nth-child(even) { background: #f5f7fa; }
			.bl-table tfoot td { background: #eef3f8; font-weight: bold; font-size: 8px; }
			.bl-center { text-align: center; }
			.bl-right { text-align: right; }
			.bl-empty { background: #fff; height: 13px; line-height: 13px; padding: 1px 3px; }

			/* ---------- TOTAL ---------- */
			.bl-total { width: 100%; border-collapse: collapse; margin-top: 6px; }
			.bl-total-label { background: #003a70; color: #fff; padding: 4px;
							font-weight: bold; text-align: right; font-size: 8px; }
			.bl-total-value { background: #003a70; color: #fff; padding: 4px;
							font-weight: bold; text-align: right; font-size: 8px; width: 30%; }

			/* ---------- SIGNATURES (compact) ---------- */
			.bl-signatures { margin-top: 8px; page-break-inside: avoid; }
			.bl-sign-table { width: 100%; border-collapse: separate; border-spacing: 6px 0; }
			.bl-sign-box { width: 50%; border: 1px solid #bbb; padding: 4px 6px;
						vertical-align: top; background: #fafbfd; }
			.bl-sign-title { margin: 0 0 3px 0; font-weight: bold; font-size: 8px;
							color: #003a70; text-transform: uppercase;
							border-bottom: 1px solid #003a70; padding-bottom: 1px; }
			.bl-sign-sub { margin: 1px 0; font-size: 8px; }
			.bl-sign-area { margin-top: 4px; height: 40px; border: 1px dashed #aaa;
							background: #fff; position: relative; }
			.bl-sign-label { position: absolute; bottom: 1px; left: 3px;
							font-size: 8px; color: #888; font-style: italic; }
			.bl-sign-mention { margin-top: 6px; font-size: 8px; font-style: italic;
							color: #555; text-align: center; }

			/* ---------- LISTE DE COLISAGE (compact) ---------- */
			.bl-lc-title { color: #003a70; font-size: 9px; font-weight: bold;
						text-align: center; margin: 6px 0 2px 0; text-transform: uppercase; }
			.bl-lc-table th { background: #003a70; color: #fff; font-size: 8px;
							text-align: center; padding: 1px 3px; }
			.bl-lc-table td { height: 13px; padding: 1px 3px; font-size: 8px; }
		</style>

		<div class="bl-container">
			<!-- ============================================================
			     TABLE UNIQUE : <thead> répété sur chaque page
			     ============================================================ -->
			<table class="bl-main-table">

				<!-- ============ EN-TÊTE (répété à chaque page) ============ -->
				<thead>
					<tr>
						<td>
							<div class="bl-header">
								<table>
									<tr>
										<td style="vertical-align: top;">
											<h1>BON DE LIVRAISON</h1>
											<p class="bl-num">N° {{ expedition.name }}</p>
										</td>
										<td class="bl-right-block" style="vertical-align: top;">
											<p><strong>Affaire :</strong> {{ expedition.project or 'N/A' }}</p>
											<p><strong>Date :</strong> {{ expedition.expedition_date or 'N/A' }}</p>
										</td>
									</tr>
								</table>
							</div>

							<table class="bl-parties">
								<tr>
									<td class="bl-party">
										<h3>EXPÉDITEUR</h3>
										<p><strong>SPIE TURBOMACHINERY</strong></p>
										<p>Z.I du Pont Long</p>
										<p>5 avenue des Frères Wright</p>
										<p>64140 LONS, France</p>
										<p>Mail : contact.tm@spie.com</p>
									</td>
									<td class="bl-party">
										<h3>DESTINATAIRE</h3>
										<p><strong>{{ expedition.client or 'N/A' }}</strong></p>
										<p>{{ data.project.location or 'N/A' }}</p>
										<p><strong>Contact :</strong> {{ data.contact }}</p>
									</td>
								</tr>
							</table>

							<table class="bl-transport">
								<tr>
									<td><strong>Transporteur :</strong> {{ carrier or 'N/A' }}</td>
									<td><strong>Date d'expédition :</strong> {{ expedition.expedition_date or 'N/A' }}</td>
									<td><strong>Poids net :</strong> {{ data.total_weight }} kg</td>
								</tr>
							</table>
						</td>
					</tr>
				</thead>

				<!-- ============ CONTENU ============ -->
				<tbody>
					<tr>
						<td>

							<!-- LISTE DE COLISAGE -->
							<table class="bl-table bl-lc-table">
								<thead>
									<tr>
										<th style="width: 12%;">Quantité de colis</th>
										<th style="width: 43%;">Désignation</th>
										<th style="width: 15%;">Poids (kg)</th>
										<th style="width: 30%;">Dimensions (cm)</th>
									</tr>
								</thead>
								<tbody>
									{% for _ in range(15) %}
									<tr>
										<td class="bl-empty">&nbsp;</td>
										<td class="bl-empty">&nbsp;</td>
										<td class="bl-empty">&nbsp;</td>
										<td class="bl-empty">&nbsp;</td>
									</tr>
									{% endfor %}
								</tbody>
								<tfoot>
									<tr>
										<td colspan="2" class="bl-right"><strong>TOTAL</strong></td>
										<td class="bl-empty">&nbsp;</td>
										<td class="bl-empty">&nbsp;</td>
									</tr>
								</tfoot>
							</table>

							<!-- COMPOSITIONS -->
							<div class="bl-lc-title">LISTE DE COLISAGE</div>
							{% if data.bom_rows and data.bom_rows|length > 0 %}
							<h3 class="bl-section">Compositions (BoM)</h3>
							<table class="bl-table">
								<thead>
									<tr>
										<th style="width: 44%;">Composition</th>
										<th style="width: 14%;">Quantité</th>
										<th style="width: 20%;">Poids unit. (kg)</th>
										<th style="width: 22%;">Nombre de colis</th>
									</tr>
								</thead>
								<tbody>
									{% for row in data.bom_rows %}
									<tr>
										<td>{{ row.composition or '' }}</td>
										<td class="bl-center">{{ row.quantity or 0 }}</td>
										<td class="bl-right">{{ '%.3f'|format(row.unit_weight) }}</td>
										<td class="bl-empty">&nbsp;</td>
									</tr>
									{% endfor %}
								</tbody>
								<tfoot>
									<tr>
										<td colspan="2" class="bl-right"><strong>Poids total compositions</strong></td>
										<td class="bl-right"><strong>{{ '%.3f'|format(data.total_weight_bom) }}</strong></td>
										<td class="bl-empty">&nbsp;</td>
									</tr>
								</tfoot>
							</table>
							{% endif %}

							<!-- ARTICLES -->
							{% if data.items_rows and data.items_rows|length > 0 %}
							<h3 class="bl-section">Articles</h3>
							<table class="bl-table">
								<thead>
									<tr>
										<th style="width: 15%;">Article</th>
										<th style="width: 37%;">Désignation</th>
										<th style="width: 12%;">Quantité</th>
										<th style="width: 16%;">Poids unit. (kg)</th>
										<th style="width: 20%;">Nombre de colis</th>
									</tr>
								</thead>
								<tbody>
									{% for row in data.items_rows %}
									<tr>
										<td>{{ row.item_from_stock or '' }}</td>
										<td>{{ row.designation or '' }}</td>
										<td class="bl-center">{{ row.true_quantity or 0 }}</td>
										<td class="bl-right">{{ '%.3f'|format(row.unit_weight) }}</td>
										<td class="bl-empty">&nbsp;</td>
									</tr>
									{% endfor %}
								</tbody>
								<tfoot>
									<tr>
										<td colspan="3" class="bl-right"><strong>Poids total articles</strong></td>
										<td class="bl-right"><strong>{{ '%.3f'|format(data.total_weight_items) }}</strong></td>
										<td class="bl-empty">&nbsp;</td>
									</tr>
								</tfoot>
							</table>
							{% endif %}

							<!-- TOTAL -->
							<table class="bl-total">
								<tr>
									<td class="bl-total-label">POIDS NET DE L'EXPÉDITION</td>
									<td class="bl-total-value">{{ '%.3f'|format(data.total_weight) }} kg</td>
								</tr>
							</table>

							<!-- SIGNATURES -->
							<div class="bl-signatures">
								<table class="bl-sign-table">
									<tr>
										<td class="bl-sign-box">
											<p class="bl-sign-title">Agent de réception / expédition</p>
											<p class="bl-sign-sub">Nom : ______________________________</p>
											<p class="bl-sign-sub">Date : ______________________________</p>
											<div class="bl-sign-area">
												<span class="bl-sign-label">Signature</span>
											</div>
										</td>
										<td class="bl-sign-box">
											<p class="bl-sign-title">Chauffeur</p>
											<p class="bl-sign-sub">Nom : ______________________________</p>
											<p class="bl-sign-sub">Date : ______________________________</p>
											<div class="bl-sign-area">
												<span class="bl-sign-label">Signature</span>
											</div>
										</td>
									</tr>
								</table>

								<p class="bl-sign-mention">
									Le signataire reconnaît avoir reçu les articles listés ci-dessus en bon état apparent.
									Toute réserve doit être formulée par écrit dans les 48 heures.
								</p>
							</div>

						</td>
					</tr>
				</tbody>
			</table>
		</div>
		""",
		{
			"expedition": expedition,
			"data": data,
			"carrier": carrier,
		},
	)

	pdf_file = get_pdf(
		html_content,
		options={
			"page-size": "A4",
			"margin-top": "8mm",
			"margin-bottom": "8mm",
			"margin-left": "8mm",
			"margin-right": "8mm",
			"encoding": "UTF-8",
		},
	)

	frappe.local.response.filename = f"BL_{expedition.name}.pdf"
	frappe.local.response.filecontent = pdf_file
	frappe.local.response.type = "pdf"


# ============================================================
# EXPORT EXCEL
# ============================================================
@frappe.whitelist()
def print_excel(name):
	"""Génère un fichier Excel (.xlsx) récapitulatif de l'expédition."""
	import openpyxl
	from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
	from io import BytesIO

	expedition = frappe.get_doc("Expedition", name)
	carrier = getattr(expedition, "carrier", None)
	data = _collect_expedition_data(expedition)

	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = "Bon de livraison"

	# Styles
	header_font = Font(bold=True, color="FFFFFF", size=10)
	header_fill = PatternFill(start_color="003A70", end_color="003A70", fill_type="solid")
	title_font = Font(bold=True, size=11, color="003A70")
	center = Alignment(horizontal="center", vertical="center")
	right = Alignment(horizontal="right", vertical="center")
	thin = Side(style="thin", color="BBBBBB")
	border = Border(left=thin, right=thin, top=thin, bottom=thin)

	# En-tête
	ws["A1"] = "BON DE LIVRAISON"
	ws["A1"].font = Font(bold=True, size=13, color="003A70")
	ws["A2"] = f"N° {expedition.name}"
	ws["A3"] = f"Affaire : {expedition.project or 'N/A'}"
	ws["A4"] = f"Date : {expedition.expedition_date or 'N/A'}"
	ws["A5"] = f"Transporteur : {carrier or 'N/A'}"
	ws["A6"] = f"Poids net : {data['total_weight']} kg"
	ws["A6"].font = Font(bold=True, size=11)

	row_idx = 8

	# Compositions
	if data["bom_rows"]:
		ws.cell(row=row_idx, column=1, value="Compositions (BoM)").font = title_font
		row_idx += 1

		headers = ["Composition", "Quantité", "Poids unit. (kg)", "Nombre de colis"]
		for col, h in enumerate(headers, start=1):
			c = ws.cell(row=row_idx, column=col, value=h)
			c.font = header_font
			c.fill = header_fill
			c.alignment = center
			c.border = border
		row_idx += 1

		for r in data["bom_rows"]:
			ws.cell(row=row_idx, column=1, value=r["composition"]).border = border
			ws.cell(row=row_idx, column=2, value=r["quantity"]).border = border
			c3 = ws.cell(row=row_idx, column=3, value=r["unit_weight"])
			c3.border = border
			c3.alignment = right
			ws.cell(row=row_idx, column=4, value="").border = border
			row_idx += 1

		ws.cell(row=row_idx, column=3, value="Poids total compositions").font = Font(bold=True)
		ws.cell(row=row_idx, column=4, value=data["total_weight_bom"]).font = Font(bold=True)
		row_idx += 2

	# Articles
	if data["items_rows"]:
		ws.cell(row=row_idx, column=1, value="Articles").font = title_font
		row_idx += 1

		headers = ["Article", "Désignation", "Quantité", "Poids unit. (kg)", "Nombre de colis"]
		for col, h in enumerate(headers, start=1):
			c = ws.cell(row=row_idx, column=col, value=h)
			c.font = header_font
			c.fill = header_fill
			c.alignment = center
			c.border = border
		row_idx += 1

		for r in data["items_rows"]:
			ws.cell(row=row_idx, column=1, value=r["item_from_stock"]).border = border
			ws.cell(row=row_idx, column=2, value=r["designation"]).border = border
			ws.cell(row=row_idx, column=3, value=r["true_quantity"]).border = border
			c4 = ws.cell(row=row_idx, column=4, value=r["unit_weight"])
			c4.border = border
			c4.alignment = right
			ws.cell(row=row_idx, column=5, value="").border = border
			row_idx += 1

		ws.cell(row=row_idx, column=4, value="Poids total articles").font = Font(bold=True)
		ws.cell(row=row_idx, column=5, value=data["total_weight_items"]).font = Font(bold=True)
		row_idx += 2

	# Total
	ws.cell(row=row_idx, column=1, value="POIDS NET DE L'EXPÉDITION").font = Font(
		bold=True, size=11, color="FFFFFF"
	)
	ws.cell(row=row_idx, column=1).fill = header_fill
	ws.cell(row=row_idx, column=2, value=f"{data['total_weight']} kg").font = Font(
		bold=True, size=11, color="FFFFFF"
	)
	ws.cell(row=row_idx, column=2).fill = header_fill

	# Largeurs de colonnes
	for col_letter, width in zip("ABCDE", [25, 45, 12, 18, 18]):
		ws.column_dimensions[col_letter].width = width

	# Retour du fichier
	stream = BytesIO()
	wb.save(stream)
	stream.seek(0)

	frappe.local.response.filename = f"BL_{expedition.name}.xlsx"
	frappe.local.response.filecontent = stream.read()
	frappe.local.response.type = "binary"
	frappe.local.response.content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
